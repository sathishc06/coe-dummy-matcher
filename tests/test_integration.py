import collections
import os
import zipfile

import openpyxl
import pytest

from conftest import subset_zip
from coe_matcher.corrections import Override, apply_overrides
from coe_matcher.pipeline import reprocess, run_reconciliation


def by_code(run):
    d = collections.defaultdict(collections.Counter)
    for r in run.results:
        d[r.subject_codes][r.status] += 1
    return {k: dict(v) for k, v in d.items()}


def test_expected_statuses_per_subject(run, dataset):
    assert by_code(run) == dataset[3]


def test_nothing_crashes_and_problem_files_are_reported(run):
    assert run.lock_files == ["DEPT_A/~$lock.xlsx"]
    assert "corrupt.xlsx" in " ".join(run.wb_problems) and "SUMMARY_ONLY" in " ".join(run.wb_problems)
    assert any(f.status == "unreadable" for f in run.catalog.files)


def test_excluded_subject_not_modified_or_written(run):
    assert not any("225PST02" in k for k in run.written)
    assert all(r.reg_no == "" for r in run.results if r.status == "EXCLUDED")


def test_no_cross_subject_assignment(run):
    row = [r for r in run.results if r.dummy_key == "1000001" and "GEOMETRY" in r.wb_rel][0]
    assert row.status == "NOT_FOUND" and row.reg_no == "" and "WRONG_SUBJECT_ONLY" in row.flags


def test_corrupt_scientific_reg_never_used_when_good_copy_exists(run):
    r = [x for x in run.results if x.dummy_key == "4000003"][0]
    assert r.status == "VERIFIED" and "E+" not in r.reg_no and r.reg_no.startswith("AC24DDD003")


def test_leading_zero_and_float_dummy_formats_match(run):
    got = {x.dummy_key: x.status for x in run.results if "STATISTICS" in x.wb_rel}
    assert got["4000001"] == got["4000002"] == "VERIFIED"


def test_ambiguous_never_gets_registration(run):
    assert all(r.reg_no == "" for r in run.results if r.status in ("AMBIGUOUS", "NOT_FOUND"))


def _sheet(path):
    return openpyxl.load_workbook(path)["Sheet1"]


def test_output_layout_sort_and_marks_association(run):
    rel = [k for k in run.written if "ALGEBRA" in k][0]
    ws = _sheet(run.written[rel])
    assert [ws.cell(7, c).value for c in (1, 2, 3, 4)] == ["Dummy No.", "Reg.No", "Reg Edit", "Marks \nScored"]
    regs = [ws.cell(r, 2).value for r in range(8, 14)]
    assert regs == sorted(regs)
    orig = _sheet(run.val_root + "/" + rel)
    om = {orig.cell(r, 1).value: tuple(orig.cell(r, c).value for c in (3, 4, 5)) for r in range(8, 14)}
    for r in range(8, 14):
        d = ws.cell(r, 1).value
        assert tuple(ws.cell(r, c).value for c in (5, 6, 7)) == om[d]                  # marks stayed with their dummy
        assert ws.cell(r, 4).value == f"=SUM(E{r}:G{r})"                               # formula follows the row
        assert ws.cell(r, 3).value == ws.cell(r, 2).value.split("-", 1)[1]


def test_unmatched_rows_listed_last_and_marked(run):
    rel = [k for k in run.written if "GEOMETRY" in k][0]
    ws = _sheet(run.written[rel])
    assert ws.cell(11, 2).value is None and ws.cell(12, 2).value is None and ws.cell(10, 2).value


def test_every_written_workbook_passed_independent_verification(run):
    assert run.verification and all(v["passed"] for v in run.verification.values())


def test_release_gate_blocks_official_zip(run):
    assert run.gate["passed"] is False and "official_zip" not in run.outputs
    assert os.path.exists(run.outputs["audit_xlsx"]) and os.path.exists(run.outputs["csv"])
    ws = openpyxl.load_workbook(run.outputs["audit_xlsx"])["Summary"]
    assert "NOT APPROVED" in str(ws["B2"].value)


def test_audit_workbook_has_required_sheets(run):
    names = openpyxl.load_workbook(run.outputs["audit_xlsx"], read_only=True).sheetnames
    for n in ["Summary", "Verified Matches", "Alternate Source Matches", "Unmatched Records", "Ambiguous Records",
              "Subject Index Crosswalk", "Workbook Integrity Verification", "Excluded Subjects"]:
        assert n in names


def test_gate_passes_and_official_zip_created_for_clean_batch(dataset, tmp_path):
    vz, dz, idx, _, _ = dataset
    clean = subset_zip(vz, str(tmp_path / "v.zip"), ["111AAA01-ALGEBRA", "225PST02"])
    r = run_reconciliation(clean, dz, idx, str(tmp_path / "w"), processes=1)
    assert r.gate["passed"] and "official_zip" in r.outputs
    assert zipfile.ZipFile(r.outputs["official_zip"]).namelist() == ["DEPT_A/111AAA01-ALGEBRA.xlsx"]


def test_manual_correction_logged_and_requires_evidence(dataset, tmp_path):
    vz, dz, idx, _, _ = dataset
    r = run_reconciliation(subset_zip(vz, str(tmp_path / "v.zip"), ["111AAA02-GEOMETRY"]), dz, idx, str(tmp_path / "w"), processes=1)
    row = [x for x in r.results if x.dummy_key == "1000001"][0]
    cand = [c for c in r.catalog.by_dummy["1000001"]][0]
    bad = Override(row.unit_key, row.sheet_row, cand.row.source_file, cand.row.row_no, "", "", "")
    assert apply_overrides(r, [bad]) == [] and row.status == "NOT_FOUND"              # incomplete evidence refused
    good = Override(row.unit_key, row.sheet_row, cand.row.source_file, cand.row.row_no, "student sat GEOMETRY under ALGEBRA bundle",
                    "attendance sheet p3", "approved by COE", user="tester")
    r2 = reprocess(r, 1, overrides=[good])
    row = [x for x in r2.results if x.dummy_key == "1000001"][0]
    assert row.status == "VERIFIED_ALTERNATE_SOURCE" and row.manual_override and row.reg_no
    log = r2.manual_log[0]
    assert log["previous_status"] == "NOT_FOUND" and log["corrected_value"] == row.reg_no and log["approval_note"]
    # original source untouched
    assert os.path.exists(os.path.join(r2.day_root, cand.row.source_file))


def test_duplicate_subject_code_conflict_is_not_silently_verified(run):
    rows = [r for r in run.results if "TOPOLOGY" in r.wb_rel]
    assert all(r.status == "VERIFIED_ALTERNATE_SOURCE" and "SUBJECT_ID_CONFLICT_RESOLVED" in r.flags for r in rows)


def test_csv_and_audit_leave_blank_cells_blank(run):
    import csv
    rows = list(csv.DictReader(open(run.outputs["csv"], encoding="utf-8-sig")))
    un = [r for r in rows if r["Status"] in ("NOT_FOUND", "AMBIGUOUS", "EXCLUDED")]
    assert un and all(r["Reg.No (final)"] == "" and r["Reg Edit"] == "" for r in un)
    ws = openpyxl.load_workbook(run.outputs["audit_xlsx"], read_only=True)["Unmatched Records"]
    first = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    assert first[4] in (None, "")
