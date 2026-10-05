"""Builds a small synthetic dataset with DELIBERATE problems. Returns (val_zip, day_zip, index_xlsx, expected)."""
import csv
import io
import os
import zipfile

import openpyxl

HDR = ["Generated_Date", "Subject_Name", "Subject_Code", "Version", "Sequence", "Bundle_No", "Registration_Number", "Dummy_ID"]


def reg(prefix, i):
    return f"{prefix}{i:03d}-24036176{i:08d}-STUDENT {i}"


def make_valuation(path, code_in_name, name, dummies, marks=None, formula_rows=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "COLLEGE"; ws.merge_cells("A1:R1")
    ws["A5"], ws["C5"] = "Programme / Branch:", "B.E/TEST"
    ws["A6"], ws["C6"] = "Sub. Code & Sub. Name:", f"{code_in_name}&{name}"
    ws.merge_cells("A5:B5"); ws.merge_cells("A6:B6")
    for j, h in enumerate(["Dummy No.", "Marks \nScored", "Q1", "Q2", "Q3"], 1):
        ws.cell(7, j, h)
    for k, d in enumerate(dummies):
        r = 8 + k
        ws.cell(r, 1, d)
        ws.cell(r, 2, f"=SUM(C{r}:E{r})")
        m = (marks[k] if marks else (k % 3 + 1, (k * 2) % 3, (k + 1) % 4))
        for j, v in enumerate(m):
            ws.cell(r, 3 + j, v)
    ws.cell(8 + len(dummies), 2, f"=SUM(C{8 + len(dummies)}:E{8 + len(dummies)})")   # blank template formula row
    wb.save(path)


def write_csv(path, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(HDR)
        w.writerows(rows)


def build(outdir):
    os.makedirs(outdir, exist_ok=True)
    V, D = os.path.join(outdir, "val"), os.path.join(outdir, "day")
    for p in (os.path.join(V, "DEPT_A"), os.path.join(D, "REGULAR", "4.5.26 FN"), os.path.join(D, "REGULAR", "5.5.26 AN")):
        os.makedirs(p, exist_ok=True)
    # ---- source files ----------------------------------------------------------------------------
    gen = "04/05/2026, 10:00:00"
    # ALGEBRA 111AAA01 : 6 students, dummies 1000001..1000006
    write_csv(os.path.join(D, "REGULAR", "4.5.26 FN", "111AAA01_ALGEBRA_2026-05-04.csv"),
              [[gen, "ALGEBRA", "111AAA01", 1, i, 1, reg("AC24AAA", i), 1000000 + i] for i in range(1, 7)])
    # GEOMETRY 111AAA02 : dummies 2000001..2000004 only (2000005 missing from source)
    write_csv(os.path.join(D, "REGULAR", "4.5.26 FN", "111AAA02_GEOMETRY_2026-05-04.csv"),
              [[gen, "GEOMETRY", "111AAA02", 1, i, 1, reg("AC24BBB", i), 2000000 + i] for i in range(1, 5)])
    # CALCULUS 111AAA03 : dummies 3000001..3000003
    write_csv(os.path.join(D, "REGULAR", "5.5.26 AN", "111AAA03_CALCULUS_2026-05-05.csv"),
              [["05/05/2026, 10:00:00", "CALCULUS", "111AAA03", 1, i, 1, reg("AC24CCC", i), 3000000 + i] for i in range(1, 4)])
    # STATISTICS 111AAA04 : xlsx with float dummies + leading zero text, one sci-notation reg + good alternate csv copy
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Sheet1"
    ws.append(HDR)
    ws.append([gen, "STATISTICS", "111AAA04", 1, 1, 1, reg("AC24DDD", 1), 4000001.0])
    ws.append([gen, "STATISTICS", "111AAA04", 1, 2, 1, reg("AC24DDD", 2), "04000002"])
    ws.append([gen, "STATISTICS", "111AAA04", 1, 3, 1, "2.40361E+15", 4000003])          # corrupted reg in this copy
    ws.append([gen, "STATISTICS", "111AAA04", 1, 4, 1, reg("AC24DDD", 4), 4000004])      # conflicts with csv copy below
    wb.save(os.path.join(D, "REGULAR", "4.5.26 FN", "111AAA04_STATISTICS_2026-05-04.xlsx"))
    write_csv(os.path.join(D, "REGULAR", "4.5.26 FN", "111AAA04_STATISTICS_copy.csv"),
              [[gen, "STATISTICS", "111AAA04", 1, 3, 1, reg("AC24DDD", 3), 4000003],
               [gen, "STATISTICS", "111AAA04", 1, 4, 1, "AC24DDD999-9999999999999999-OTHER", 4000004]])
    # conflicting-subject workbook 111AAA06 (file name says 111AAA05)
    write_csv(os.path.join(D, "REGULAR", "5.5.26 AN", "111AAA06_TOPOLOGY_2026-05-05.csv"),
              [["05/05/2026, 10:00:00", "TOPOLOGY", "111AAA06", 1, i, 1, reg("AC24FFF", i), 6000000 + i] for i in range(1, 4)])
    # excluded subject
    write_csv(os.path.join(D, "REGULAR", "4.5.26 FN", "225PST02_PS_2026-05-04.csv"),
              [[gen, "P S", "225PST02", 1, i, 1, reg("AC24PPP", i), 9000000 + i] for i in range(1, 3)])
    # corrupt + lock files in Daywise
    open(os.path.join(D, "REGULAR", "4.5.26 FN", "broken.xlsx"), "wb").write(b"this is not an excel file")
    open(os.path.join(D, "REGULAR", "4.5.26 FN", "~$lock.xlsx"), "wb").write(b"lock")
    # ---- official index ----------------------------------------------------------------------------
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Sheet1"
    ws.append(["TEST INDEX"]); ws.append([None]); ws.append([])
    ws.append(["S.No", "Exam Date", "Day", "Session", "Sem", "Sub. Code", "Sub. Name", "Dept", "REG", "ARR", "OG/PHD", "Daywise Strength"])
    ents = [("04.05.2026", "Monday", "FN", "111AAA01", "ALGEBRA", 6), ("04.05.2026", "Monday", "FN", "111AAA02", "GEOMETRY", 5),
            ("05.05.2026", "Tuesday", "AN", "111AAA03", "CALCULUS", 3), ("04.05.2026", "Monday", "FN", "111AAA04", "STATISTICS", 4),
            ("05.05.2026", "Tuesday", "AN", "111AAA05", "TOPOLOGY OLD", 3), ("05.05.2026", "Tuesday", "AN", "111AAA06", "TOPOLOGY", 3),
            ("04.05.2026", "Monday", "FN", "225PST02", "P S", 2)]
    for i, (d, day, s, c, n, k) in enumerate(ents, 1):
        ws.append([i, d, day, s, 4, c, n, "TESTDEPT", k, 0, 0, k])
    idx = os.path.join(outdir, "INDEX.xlsx")
    wb.save(idx)
    # ---- valuation workbooks -------------------------------------------------------------------------
    A = os.path.join(V, "DEPT_A")
    # unsorted dummies on purpose (reg order != dummy order)
    make_valuation(os.path.join(A, "111AAA01-ALGEBRA.xlsx"), "111AAA01", "ALGEBRA", [1000003, 1000001, 1000006, 1000002, 1000005, 1000004])
    make_valuation(os.path.join(A, "111AAA02-GEOMETRY.xlsx"), "111AAA02", "GEOMETRY", [2000001, 2000002, 1000001, 2000005, 2000003])
    make_valuation(os.path.join(A, "111AAA03-CALCULUS.xlsx"), "111AAA03", "CALCULUS", [3000001, 3000002, 3000002, 3000003],
                   marks=[(1, 1, 1), (2, 2, 2), (0, 1, 0), (3, 3, 3)])
    make_valuation(os.path.join(A, "111AAA04-STATISTICS.xlsx"), "111AAA04", "STATISTICS", [4000001, 4000002, 4000003, 4000004])
    make_valuation(os.path.join(A, "111AAA05-TOPOLOGY.xlsx"), "111AAA06", "TOPOLOGY", [6000001, 6000002, 6000003])
    make_valuation(os.path.join(A, "225PST02-PS.xlsx"), "225PST02", "P S", [9000001, 9000002])
    open(os.path.join(A, "~$lock.xlsx"), "wb").write(b"lock")
    open(os.path.join(A, "corrupt.xlsx"), "wb").write(b"not excel")
    wb = openpyxl.Workbook(); wb.active["A1"] = "just a summary sheet"; wb.save(os.path.join(A, "SUMMARY_ONLY.xlsx"))

    def zipdir(src, dst):
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
            for dp, _, fns in os.walk(src):
                for fn in fns:
                    z.write(os.path.join(dp, fn), os.path.relpath(os.path.join(dp, fn), src))
        return dst
    vz, dz = zipdir(V, os.path.join(outdir, "VAL.zip")), zipdir(D, os.path.join(outdir, "DAY.zip"))
    expected = {
        "111AAA01": {"VERIFIED": 6},
        "111AAA02": {"VERIFIED": 3, "NOT_FOUND": 2},           # missing record + dummy that belongs to ALGEBRA only
        "111AAA03": {"VERIFIED": 2, "AMBIGUOUS": 2},           # duplicate dummy with different marks
        "111AAA04": {"VERIFIED": 3, "AMBIGUOUS": 1},           # 4000004 conflicts across two copies
        "111AAA06": {"VERIFIED_ALTERNATE_SOURCE": 3},          # file name 05 vs content 06 conflict, resolved by name equality
        "225PST02": {"EXCLUDED": 2},
    }
    return vz, dz, idx, expected


if __name__ == "__main__":
    import sys
    print(build(sys.argv[1] if len(sys.argv) > 1 else "testdata"))
