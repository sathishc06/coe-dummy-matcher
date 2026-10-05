"""Official date/session index (DAYWISE workbook) parser."""
import datetime as dt
import re
from dataclasses import dataclass, field
from typing import List, Optional

import openpyxl

from .normalize import clean_text, name_key

DAY_NAMES = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
TOKEN_SPLIT = re.compile(r"[\s/,&;]+")
OFFICIAL_CODE = re.compile(r"^\d{3}[A-Z]{3}\d{2}$")


@dataclass
class IndexEntry:
    row_no: int
    sno: object
    date_raw: str
    date: Optional[str]            # ISO yyyy-mm-dd actually used
    date_note: str                 # '' or explanation of a correction
    day: str
    session: str                   # FN | AN | ''
    sem: object
    code_raw: str
    tokens: List[str]
    malformed_tokens: List[str]
    name: str
    dept: str
    reg: Optional[int]
    arr: Optional[int]
    og: Optional[int]
    strength: Optional[int]
    qp: str = ""
    sheet: str = ""

    @property
    def name_key(self):
        return name_key(self.name)


@dataclass
class IndexData:
    path: str
    sheet: str = ""
    entries: List[IndexEntry] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    found: bool = False


def _num(v):
    try:
        if v is None or v == "":
            return None
        return int(float(v))
    except Exception:
        return None


def parse_date(s):
    s = clean_text(s)
    m = re.match(r"^(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})$", s)
    if not m:
        return None
    d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if y < 100:
        y += 2000
    try:
        return dt.date(y, mo, d)
    except ValueError:
        return None


def norm_session(s):
    t = re.sub(r"[^A-Z]", "", clean_text(s).upper())
    if t in ("FN", "F", "FORENOON", "MORNING"):
        return "FN"
    if t in ("AN", "A", "AFTERNOON"):
        return "AN"
    return ""


def split_tokens(code_raw):
    toks = [t.upper() for t in TOKEN_SPLIT.split(clean_text(code_raw)) if t]
    good = [t for t in toks if OFFICIAL_CODE.match(t)]
    bad = [t for t in toks if not OFFICIAL_CODE.match(t)]
    return good, bad


def find_index_sheet(wb):
    """Find the sheet + header row that looks like the official date/session index."""
    for ws in wb.worksheets:
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), 1):
            hs = {re.sub(r"[^a-z]", "", clean_text(v).lower()): j for j, v in enumerate(row) if v is not None}
            if "subcode" in hs and "examdate" in hs and "session" in hs:
                return ws, i, hs
    return None, None, None


def load_index(path) -> IndexData:
    data = IndexData(path)
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        data.warnings.append(f"cannot open index workbook: {e!r}")
        return data
    ws, hdr_row, hs = find_index_sheet(wb)
    if ws is None:
        data.warnings.append("no sheet with 'Exam Date' / 'Session' / 'Sub. Code' headers was found")
        return data
    data.found, data.sheet = True, ws.title
    col = lambda *names: next((hs[n] for n in names if n in hs), None)
    c_date, c_day, c_sess = col("examdate"), col("day"), col("session")
    c_sem, c_code, c_name = col("sem"), col("subcode"), col("subname")
    c_reg, c_arr, c_og = col("reg"), col("arr"), col("ogphd")
    c_dept, c_str, c_qp = col("dept"), col("daywisestrength", "totalnoofstudents"), col("noofqpsprinted")
    years = []
    raw_rows = []
    for i, row in enumerate(ws.iter_rows(min_row=hdr_row + 1, values_only=True), hdr_row + 1):
        g = lambda c: row[c] if c is not None and c < len(row) else None
        if clean_text(g(c_code)) == "" and clean_text(g(c_date)) == "":
            if any(v is not None for v in row):
                data.warnings.append(f"index row {i}: incomplete row skipped ({[v for v in row if v is not None][:3]})")
            continue
        d = parse_date(g(c_date))
        if d:
            years.append(d.year)
        raw_rows.append((i, row, d))
    main_year = max(set(years), key=years.count) if years else None
    for i, row, d in raw_rows:
        g = lambda c: row[c] if c is not None and c < len(row) else None
        note = ""
        day = clean_text(g(c_day))
        if d is not None:
            wd = DAY_NAMES[d.weekday()]
            if day and wd.lower() != day.lower():
                alt = None
                if main_year and d.year != main_year:
                    try:
                        alt = d.replace(year=main_year)
                    except ValueError:
                        alt = None
                if alt and DAY_NAMES[alt.weekday()].lower() == day.lower():
                    note = (f"date '{clean_text(g(c_date))}' disagrees with Day='{day}'; year corrected to {main_year} "
                            f"because that matches the weekday")
                    d = alt
                else:
                    note = f"date '{clean_text(g(c_date))}' disagrees with Day='{day}' (weekday check failed)"
        code_raw = clean_text(g(c_code))
        good, bad = split_tokens(code_raw)
        e = IndexEntry(
            row_no=i, sno=g(0), date_raw=clean_text(g(c_date)), date=d.isoformat() if d else None, date_note=note,
            day=day, session=norm_session(g(c_sess)), sem=g(c_sem), code_raw=code_raw, tokens=good,
            malformed_tokens=bad, name=clean_text(g(c_name)), dept=clean_text(g(c_dept)),
            reg=_num(g(c_reg)), arr=_num(g(c_arr)), og=_num(g(c_og)), strength=_num(g(c_str)),
            qp=clean_text(g(c_qp)), sheet=ws.title)
        if bad:
            data.warnings.append(f"index row {i}: malformed code token(s) {bad} in '{code_raw}'")
        if note:
            data.warnings.append(f"index row {i}: {note}")
        data.entries.append(e)
    return data
