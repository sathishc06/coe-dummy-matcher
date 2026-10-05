"""Audit workbook, CSV, text report and ZIP packages."""
import csv
import os
import zipfile
import re
import copy
import shutil
from collections import Counter, defaultdict

from openpyxl import Workbook
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, column_index_from_string

from .gate import REL_BLOCKED, REL_COMPLETE, REL_EXCL, REL_NA, REL_NONE, REL_PARTIAL
from .models import (AMBIGUOUS, EXCLUDED, MATCHED_STATUSES, NOT_FOUND, VERIFIED, VERIFIED_ALT)

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY = Font(name="Arial", size=10)
NOT_APPROVED = "NOT APPROVED FOR OFFICIAL USE"

ROW_COLS = ["Workbook", "Sheet", "Sheet row", "Original dummy", "Reg.No (final)", "Reg Edit", "Status", "Subject code(s)",
            "Subject name", "Exam date", "Session", "Sem", "Dept", "Source file", "Source sheet", "Source row",
            "Source subject code", "Subject evidence", "Other agreeing copies", "Verification", "Flags", "Reason / note"]


def row_values(r):
    return [r.wb_rel, r.sheet, r.sheet_row, r.dummy_raw if r.dummy_raw is not None else "", r.reg_no, r.reg_edit, r.status,
            r.subject_codes, r.subject_name, r.exam_date, r.session, r.sem, r.dept, r.src_file, r.src_sheet, r.src_row,
            r.src_subject_code, r.src_evidence, r.other_sources, r.verification, ", ".join(r.flags), r.reason]


def _sheet(wb, title, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    n = 0
    for v in rows:
        ws.append([("'" + x if isinstance(x, str) and x != "" and x[0] in "=+-@" else x) for x in v])
        n += 1
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, min(48, max(12, len(str(h)) + 4)))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, n + 1)}"
    return ws


def counts(run):
    c = Counter(r.status for r in run.results)
    return c


def _release_counts(run):
    return Counter(st for st, _ in run.wb_status.values())


def summary_numbers(run):
    c = counts(run)
    total_files = len(run.val_report.files) if run.val_report else 0
    wbs = {u.wb_rel for u in run.units} | set(run.wb_problems) | set(run.lock_files)
    rc = _release_counts(run)
    verified_wbs = [rel for rel, v in run.verification.items() if v["passed"]]
    return dict(workbooks_examined=len(wbs), lock_files=len(run.lock_files), rows=len(run.results),
                verified=c[VERIFIED], alt=c[VERIFIED_ALT], ambiguous=c[AMBIGUOUS], not_found=c[NOT_FOUND],
                excluded=c[EXCLUDED], wb_written=len(run.written), wb_verified=len(verified_wbs),
                wb_complete=rc[REL_COMPLETE], wb_partial=rc[REL_PARTIAL], wb_blocked=rc[REL_BLOCKED] + 0,
                wb_none=rc[REL_NONE], wb_excluded=rc[REL_EXCL], wb_nostud=rc[REL_NA])




def _dept_code_from_reg(reg_no, prefix="6176"):
    """Extract the first 3 consecutive digits occurring after the registration prefix.
    Example: 6176...101... -> 101. Returns UNKNOWN_DEPT when unavailable.
    """
    s = re.sub(r"\s+", "", str(reg_no or "")).upper()
    pos = s.find(prefix)
    if pos < 0:
        return "UNKNOWN_DEPT"
    tail = s[pos + len(prefix):]
    m = re.search(r"(\d{3})", tail)
    return m.group(1) if m else "UNKNOWN_DEPT"


def _primary_subject_code(row, unit=None):
    codes = re.findall(r"\d{3}[A-Z]{3}\d{2}", str(getattr(row, "subject_codes", "") or "").upper())
    if codes:
        return codes[0]
    if unit and getattr(unit, "codes", None):
        return str(unit.codes[0]).upper()
    return "UNKNOWN_SUBJECT"


def _copy_cell(src, dst, old_row=None, new_row=None):
    dst.value = src.value
    if isinstance(dst.value, str) and dst.value.startswith("=") and old_row is not None and new_row is not None:
        # Row-only formula rewrite; columns are unchanged in the already processed workbook.
        pattern = re.compile(r"(?<![A-Za-z0-9_$.!\"])(\$?[A-Z]{1,3})(\$?)(\d+)(?![A-Za-z0-9_(])")
        def sub(m):
            r = int(m.group(3))
            return f"{m.group(1)}{m.group(2)}{new_row if r == old_row else r}"
        dst.value = pattern.sub(sub, dst.value)
    if src.has_style:
        dst._style = copy.copy(src._style)
    if src.number_format:
        dst.number_format = src.number_format
    if src.alignment:
        dst.alignment = copy.copy(src.alignment)
    if src.protection:
        dst.protection = copy.copy(src.protection)
    if src.hyperlink:
        dst._hyperlink = copy.copy(src.hyperlink)
    if src.comment:
        dst.comment = copy.copy(src.comment)


def _write_subject_dept_workbook(src_path, sheet_name, header_row, student_rows, out_path, subject_code, dept_code, source_rel):
    """Create a clean subject+department Excel containing only selected verified student rows.
    Header/title rows are preserved; student rows are compacted and formulas are row-adjusted.
    """
    src_wb = openpyxl.load_workbook(src_path)
    src_ws = src_wb[sheet_name]
    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = re.sub(r"[^A-Za-z0-9_ -]", "_", subject_code)[:31] or "Subject"
    max_col = src_ws.max_column
    # Preserve column widths and hidden state.
    for key, dim in src_ws.column_dimensions.items():
        out_ws.column_dimensions[key].width = dim.width
        out_ws.column_dimensions[key].hidden = dim.hidden
    # Copy title/header area.
    for r in range(1, header_row + 1):
        out_ws.row_dimensions[r].height = src_ws.row_dimensions[r].height
        for c in range(1, max_col + 1):
            _copy_cell(src_ws.cell(r, c), out_ws.cell(r, c))
    # Compact selected student rows directly after the header.
    for new_row, old_row in enumerate(student_rows, start=header_row + 1):
        out_ws.row_dimensions[new_row].height = src_ws.row_dimensions[old_row].height
        for c in range(1, max_col + 1):
            _copy_cell(src_ws.cell(old_row, c), out_ws.cell(new_row, c), old_row, new_row)
    # Copy only merges entirely inside the title/header area. Student-row merges are intentionally
    # omitted to avoid carrying merges from excluded departments.
    for m in list(src_ws.merged_cells.ranges):
        if m.max_row <= header_row:
            out_ws.merge_cells(str(m))
    out_ws.freeze_panes = src_ws.freeze_panes if src_ws.freeze_panes and str(src_ws.freeze_panes).isdigit() is False else None
    out_ws.auto_filter.ref = f"A{header_row}:{get_column_letter(max_col)}{header_row + len(student_rows)}"
    out_ws.sheet_view.showGridLines = src_ws.sheet_view.showGridLines
    # Release metadata makes the split provenance explicit.
    meta = out_wb.create_sheet("Release Info")
    meta.append(["Field", "Value"])
    for cell in meta[1]:
        cell.font = Font(bold=True)
    meta_rows = [
        ("Subject code", subject_code), ("Department code", dept_code),
        ("Registration parsing rule", "First 3 consecutive digits after 6176 in Reg.No"),
        ("Source valuation workbook", source_rel), ("Source worksheet", sheet_name),
        ("Rows included", len(student_rows)), ("Release basis", "All included rows matched and passed independent verification"),
    ]
    for x in meta_rows:
        meta.append(list(x))
    meta.column_dimensions["A"].width = 30
    meta.column_dimensions["B"].width = 100
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    out_wb.save(out_path)


def build_subject_dept_release(run, out_dir):
    """Build the requested usable release: exclude every non-COMPLETE workbook; split output by subject and dept.
    No manual correction or second verification step is required here; only rows already classified COMPLETE_VERIFIED
    by the existing independent verification stage are eligible.
    """
    base = os.path.join(out_dir, "OFFICIAL_USABLE_SUBJECTS")
    shutil.rmtree(base, ignore_errors=True)
    os.makedirs(base, exist_ok=True)
    unit_map = {f"{u.wb_rel}::{u.sheet}": u for u in run.units}
    groups = defaultdict(list)
    for rel, (st, _) in run.wb_status.items():
        if st != REL_COMPLETE or rel not in run.written:
            continue
        for r in run.results:
            if r.wb_rel != rel or r.status not in MATCHED_STATUSES or r.verification != "PASS":
                continue
            u = unit_map.get(r.unit_key)
            if not u:
                continue
            subject = _primary_subject_code(r, u)
            dept = _dept_code_from_reg(r.reg_no)
            groups[(subject, dept, rel, r.sheet)].append(r)
    files = []
    counters = Counter()
    for (subject, dept, rel, sheet), rows in sorted(groups.items()):
        u = unit_map.get(f"{rel}::{sheet}")
        if not u:
            continue
        safe_subject = re.sub(r"[^A-Za-z0-9_-]", "_", subject)
        safe_dept = re.sub(r"[^A-Za-z0-9_-]", "_", dept)
        folder = os.path.join(base, safe_subject)
        counters[(subject, dept)] += 1
        suffix = f"_{counters[(subject, dept)]:02d}" if counters[(subject, dept)] > 1 else ""
        out_path = os.path.join(folder, f"{safe_subject}_{safe_dept}{suffix}.xlsx")
        _write_subject_dept_workbook(run.written[rel], sheet, u.header_row, [r.sheet_row for r in rows], out_path, subject, dept, rel)
        files.append((out_path, subject, dept, len(rows)))
    zip_path = os.path.join(out_dir, "USABLE_SUBJECT_DEPARTMENT_RELEASE.zip")
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("RELEASE_NOTE.txt",
                   "USABLE SUBJECT/DEPARTMENT RELEASE\n\n"
                   "Only workbooks classified COMPLETE_VERIFIED and rows independently verified as PASS are included.\n"
                   "Subjects/workbooks with unresolved, ambiguous, blocked, excluded or unverified data are omitted.\n"
                   "Department code = first 3 consecutive digits found after the 6176 prefix in the complete Reg.No.\n"
                   "No manual correction/recheck is performed during this release packaging step.\n")
        for p, subject, dept, n in files:
            z.write(p, os.path.relpath(p, base))
    return zip_path, files

def write_all_outputs(run, out_dir):
    os.makedirs(out_dir, exist_ok=True)
    o = {}
    S = summary_numbers(run)
    gate_ok = run.gate.get("passed", False)
    label = "RELEASE GATE PASSED" if gate_ok else NOT_APPROVED
    # ---------------- audit workbook ----------------
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Summary")
    lines = [
        ("FINAL RECONCILIATION AUDIT", ""), ("Release status", label),
        ("Strict release rule", "100% completion is NOT claimed unless every included row has an authoritative mapping and every output workbook passes independent verification."),
        ("", ""),
        ("Valuation workbooks examined (incl. lock files)", S["workbooks_examined"]),
        ("  Office lock files (~$) skipped", S["lock_files"]),
        ("Total student rows examined", S["rows"]),
        ("VERIFIED matches", S["verified"]), ("VERIFIED_ALTERNATE_SOURCE matches", S["alt"]),
        ("NOT_FOUND (unmatched) records", S["not_found"]), ("AMBIGUOUS records", S["ambiguous"]),
        ("EXCLUDED rows (" + ", ".join(run.excluded_codes) + ")", S["excluded"]),
        ("Rows with a final Reg.No", S["verified"] + S["alt"]),
        ("", ""),
        ("Corrected workbooks written", S["wb_written"]), ("Workbooks passing independent verification", S["wb_verified"]),
        ("Workbooks COMPLETE (every row mapped + verified)", S["wb_complete"]),
        ("Workbooks PARTIAL (exceptions remain) - blocked from official release", S["wb_partial"]),
        ("Workbooks BLOCKED (unsupported structure / failed verification)", S["wb_blocked"]),
        ("Workbooks with nothing verifiable (no output)", S["wb_none"]),
        ("Workbooks excluded", S["wb_excluded"]),
        ("", ""),
        ("Daywise source files parsed", len(run.catalog.files) if run.catalog else 0),
        ("  CSV / Excel", f"{run.catalog.n_csv} / {run.catalog.n_xlsx}" if run.catalog else ""),
        ("Daywise source records", len(run.catalog.recs) if run.catalog else 0),
        ("Official index entries", len(run.index.entries) if run.index else 0),
    ]
    for a, b in lines:
        ws.append([a, b])
    ws["A1"].font = Font(name="Arial", bold=True, size=14)
    for r in ws.iter_rows(min_row=2):
        r[0].font = Font(name="Arial", bold=True, size=10)
        r[1].font = BODY
        r[1].alignment = Alignment(wrap_text=True, vertical="top")
    if not gate_ok:
        ws["B2"].font = Font(name="Arial", bold=True, color="C00000")
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 80
    # row sheets
    _sheet(wb, "Verified Matches", ROW_COLS, (row_values(r) for r in run.results if r.status == VERIFIED))
    _sheet(wb, "Alternate Source Matches", ROW_COLS, (row_values(r) for r in run.results if r.status == VERIFIED_ALT))
    _sheet(wb, "Unmatched Records", ROW_COLS, (row_values(r) for r in run.results if r.status == NOT_FOUND))
    _sheet(wb, "Ambiguous Records", ROW_COLS, (row_values(r) for r in run.results if r.status == AMBIGUOUS))
    # crosswalk
    ch = ["Workbook", "Sheet", "Subject code(s) used", "Sheet subject name", "Identification", "Identification note",
          "Crosswalk status", "Index date(s)", "Session(s)", "Index sem", "Index dept", "Index subject name",
          "Index strength (Reg/Arr/OG)", "Student rows", "VERIFIED", "ALTERNATE", "NOT_FOUND", "AMBIGUOUS", "EXCLUDED",
          "Source files used", "Notes"]

    def cw_rows():
        for s in run.summaries:
            es = s.crosswalk.entries
            yield [s.wb_rel, s.sheet, s.codes, s.subject_name, s.id_status, s.id_note, s.crosswalk.status + (f" - {s.crosswalk.note}" if s.crosswalk.note else ""),
                   "; ".join(sorted({e.date or "" for e in es})), "; ".join(sorted({e.session for e in es})),
                   "; ".join(sorted({str(e.sem) for e in es})), "; ".join(sorted({e.dept for e in es})),
                   "; ".join(sorted({e.name for e in es})),
                   "; ".join(f"{e.reg}/{e.arr}/{e.og}" for e in es), s.n_rows,
                   s.counts.get(VERIFIED, 0), s.counts.get(VERIFIED_ALT, 0), s.counts.get(NOT_FOUND, 0),
                   s.counts.get(AMBIGUOUS, 0), s.counts.get(EXCLUDED, 0), "; ".join(s.sources_used[:5]),
                   "; ".join(s.notes + ([f"structure: {x}" for x in s.structure_flags]))]
    _sheet(wb, "Subject Index Crosswalk", ch, cw_rows())
    # integrity
    ih = ["Workbook", "Release status", "Release reason", "Written", "Independent verification", "Failed checks",
          "Checks run", "Output file"]

    def int_rows():
        for rel in sorted(run.wb_status):
            st, why = run.wb_status[rel]
            v = run.verification.get(rel)
            failed = "; ".join(f"{n}: {d}" for n, ok, d in v["checks"] if not ok) if v else ""
            yield [rel, st, why, "YES" if rel in run.written else "NO", ("PASS" if v["passed"] else "FAIL") if v else "n/a",
                   failed, len(v["checks"]) if v else 0, run.written.get(rel, "")]
    _sheet(wb, "Workbook Integrity Verification", ih, int_rows())
    # excluded
    _sheet(wb, "Excluded Subjects", ["Workbook", "Sheet", "Subject code(s)", "Student rows", "Reason"],
           ([s.wb_rel, s.sheet, s.codes, s.n_rows, "Excluded by instruction - not modified, not reconciled; manual verification by COE"]
            for s in run.summaries if s.excluded))
    # index problems
    if run.index and run.index.warnings:
        _sheet(wb, "Index Warnings", ["Warning"], ([w] for w in run.index.warnings))
    if getattr(run, "manual_log", None):
        from .corrections import LOG_FIELDS
        _sheet(wb, "Manual Corrections Log", LOG_FIELDS, ([m[k] for k in LOG_FIELDS] for m in run.manual_log))
    p = os.path.join(out_dir, "FINAL_RECONCILIATION_AUDIT.xlsx")
    wb.save(p)
    o["audit_xlsx"] = p
    # ---------------- CSV ----------------
    p = os.path.join(out_dir, "ALL_STUDENT_ROWS_STATUS.csv")
    with open(p, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(ROW_COLS)
        for r in run.results:
            w.writerow([("'" + x if isinstance(x, str) and x != "" and x[0] in "=+-@" else x) for x in row_values(r)])
    o["csv"] = p
    # ---------------- report ----------------
    p = os.path.join(out_dir, "VERIFICATION_REPORT.md")
    open(p, "w", encoding="utf-8").write(build_report(run, S, label))
    o["report"] = p
    # ---------------- exceptions CSV ----------------
    p = os.path.join(out_dir, "EXCEPTION_REPORT.csv")
    with open(p, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(ROW_COLS)
        for r in run.results:
            if r.status in (NOT_FOUND, AMBIGUOUS):
                w.writerow(row_values(r))
    o["exceptions_csv"] = p
    # ---------------- ZIPs ----------------
    base = os.path.join(out_dir, "corrected_workbooks")
    complete = [rel for rel, (st, _) in run.wb_status.items() if st == REL_COMPLETE and rel in run.written]
    partial = [rel for rel, (st, _) in run.wb_status.items() if st == REL_PARTIAL and rel in run.written]
    if gate_ok:
        p = os.path.join(out_dir, "CORRECTED_WORKBOOKS_OFFICIAL.zip")
        with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
            for rel in complete:
                z.write(run.written[rel], rel)
        o["official_zip"] = p
    p = os.path.join(out_dir, "CORRECTED_WORKBOOKS_COMPLETE_ONLY.zip")
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("README.txt", "Workbooks in this ZIP had every included student row mapped and passed independent "
                   "verification.\n" + (NOT_APPROVED + " - the overall release gate FAILED (other workbooks still have exceptions).\n" if not gate_ok else ""))
        for rel in complete:
            z.write(run.written[rel], rel)
    o["complete_zip"] = p
    # User-authorized scoped release: only COMPLETE_VERIFIED workbooks.
    # This is intentionally distinct from the strict overall gate. Any PARTIAL,
    # BLOCKED or NO_OUTPUT workbook is excluded from this ZIP.
    p = os.path.join(out_dir, "CORRECTED_WORKBOOKS_SCOPED_OFFICIAL_USE.zip")
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("README.txt",
                   "SCOPED OFFICIAL-USE RELEASE\n\n"
                   "This ZIP contains ONLY workbooks classified COMPLETE_VERIFIED: every included student row was mapped and the output workbook passed independent verification.\n"
                   "All workbooks that failed the overall strict release gate are intentionally excluded.\n"
                   "This file must NOT be interpreted as evidence that the overall May 2026 reconciliation release gate passed.\n")
        for rel in complete:
            z.write(run.written[rel], rel)
    o["scoped_official_zip"] = p
    # Requested usable release: automatically excludes unusable workbooks and splits by subject/department.
    p, release_files = build_subject_dept_release(run, out_dir)
    o["usable_subject_department_zip"] = p
    o["usable_subject_department_files"] = release_files
    p = os.path.join(out_dir, "CORRECTED_WORKBOOKS_PARTIAL_NOT_FOR_RELEASE.zip")
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("README.txt", NOT_APPROVED + "\nThese workbooks contain verified Reg.No values for matched rows only. "
                   "Unmatched rows are listed last with a yellow-highlighted blank Reg.No / Reg Edit. See the audit workbook for "
                   "the exact reason per row.\n")
        for rel in partial:
            z.write(run.written[rel], rel)
    o["partial_zip"] = p
    return o


def build_report(run, S, label):
    L = []
    A = L.append
    A("# COE Dummy-Number Reconciliation - Verification Report")
    A(f"\n**Release status: {label}**\n")
    A("| Measure | Count |\n|---|---:|")
    for k, v in [("Valuation workbooks examined", S["workbooks_examined"]), ("  of which Office lock files (~$) skipped", S["lock_files"]),
                 ("Student rows examined", S["rows"]), ("VERIFIED matches", S["verified"]),
                 ("VERIFIED_ALTERNATE_SOURCE matches", S["alt"]), ("Unmatched (NOT_FOUND) records", S["not_found"]),
                 ("Ambiguous records", S["ambiguous"]), ("Excluded rows", S["excluded"]),
                 ("Corrected workbooks written", S["wb_written"]), ("Workbooks passing independent verification", S["wb_verified"]),
                 ("Workbooks COMPLETE (releasable)", S["wb_complete"]), ("Workbooks PARTIAL (blocked from release)", S["wb_partial"]),
                 ("Workbooks BLOCKED (structure / verification)", S["wb_blocked"]),
                 ("Workbooks with nothing verifiable", S["wb_none"])]:
        A(f"| {k} | {v:,} |")
    # reasons
    nf = [r for r in run.results if r.status in (NOT_FOUND, AMBIGUOUS)]
    A("\n## Exact reasons for unresolved records (grouped)\n")
    grp = Counter()
    for r in nf:
        key = r.flags[0] if r.flags else r.status
        grp[key] += 1
    expl = {"NO_SOURCE_RECORD": "dummy number not present in any Daywise source file",
            "WRONG_SUBJECT_ONLY": "dummy number exists only in the source of a DIFFERENT subject (cross-subject matching prohibited)",
            "DUPLICATE_DUMMY_IN_WORKBOOK_DIFFERENT_MARKS": "same dummy appears on several rows with different marks",
            "SOURCE_REG_UNUSABLE": "only source copy has a corrupted/blank registration number",
            "INVALID_DUMMY_VALUE": "dummy cell is not numeric",
            "CONFLICTING_MAPPING": "conflicting registration numbers within the subject's evidence"}
    A("| Reason | Rows |\n|---|---:|")
    for k, v in grp.most_common():
        A(f"| {expl.get(k, k)} | {v:,} |")
    A("\nEvery unresolved row carries its own exact reason in `ALL_STUDENT_ROWS_STATUS.csv` / `EXCEPTION_REPORT.csv`.")
    # needed evidence
    byu = defaultdict(list)
    for r in nf:
        byu[(r.wb_rel, r.sheet)].append(r)
    A("\n## Additional source evidence needed (largest gaps)\n")
    A("| Workbook / sheet | Unresolved rows | Evidence required |\n|---|---:|---|")
    for (wbn, sh), rs in sorted(byu.items(), key=lambda kv: -len(kv[1]))[:40]:
        f = Counter(r.flags[0] if r.flags else "" for r in rs).most_common(1)[0][0]
        need = {"NO_SOURCE_RECORD": "Daywise dummy-to-register file (e.g. ARREAR/supplementary list) covering these dummy numbers",
                "WRONG_SUBJECT_ONLY": "Confirmation of which subject these dummy numbers belong to, or the correct subject's Daywise file",
                "DUPLICATE_DUMMY_IN_WORKBOOK_DIFFERENT_MARKS": "Corrected valuation sheet (duplicate dummy numbers with different marks)"}.get(f, "Authoritative source record")
        A(f"| {wbn} / {sh} | {len(rs):,} | {need} |")
    # blocked workbooks
    bl = [(rel, v) for rel, v in run.wb_status.items() if v[0] == REL_BLOCKED]
    if bl:
        A("\n## Workbooks blocked from release\n")
        for rel, (st, why) in bl:
            A(f"- `{rel}` - {why}")
    fails = [(rel, [n for n, ok, d in v['checks'] if not ok]) for rel, v in run.verification.items() if not v["passed"]]
    if fails:
        A("\n## Independent verification failures\n")
        for rel, f in fails:
            A(f"- `{rel}`: {f}")
    A("\n## Notes\n- Reg Edit is blank where the registration number contains no hyphen (text after the first hyphen does not exist).")
    A("- Workbooks are saved with *recalculate on open*; marks formulas were rewritten for the shifted columns and verified by an independent evaluator.")
    A(f"- Run timings (s): {run.timings}")
    return "\n".join(L)
